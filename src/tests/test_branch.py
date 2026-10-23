"""부점 브리핑(tools/branch_briefing) 회귀 테스트 — SQL 검사 · 부점 범위 · 행 수 상한 · 미조회 조건 · 재작성 · 원장·출처 ·
룰베이스 부점 시트 적재 · 현황·CASE 판정 · 전략 브리핑 숫자 검사 · 관리 대상 추천 순서·첫 추천 유지. LLM 스텁.

실행: python -m tests.test_branch
"""

from __future__ import annotations

import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from pension_agent import tools
from pension_agent.tools.base import ToolFailure
from pension_agent.tools.branch_briefing import check as CK
from pension_agent.tools.branch_briefing import run, tool, views

_results: list[tuple[bool, str, str]] = []

EMP = "3902172"          # 시연 사번 — 여의도종합금융센터(0001)
YEOUIDO = "여의도종합금융센터"
OTHER_EMP = "3902201"    # 강남(0002)


def check(cond: bool, label: str, detail: str = "") -> None:
    _results.append((bool(cond), label, detail))
    print(f"{'✓' if cond else '✗'} {label}" + (f" — {detail}" if detail and not cond else ""))


class _LLM:
    """tool.generate 교체 — replies 순차 반환, 프롬프트 기록."""

    def __init__(self, *replies: str):
        self.replies = list(replies)
        self.prompts: list[str] = []

    def __enter__(self) -> _LLM:
        self._saved = tool.generate

        def _gen(prompt, **kw):
            self.prompts.append(prompt)
            return self.replies.pop(0) if len(self.replies) > 1 else self.replies[0]
        tool.generate = _gen
        return self

    def __exit__(self, *exc) -> None:
        tool.generate = self._saved


def _sql(s: str | None, reason: str = "확인") -> str:
    import json
    return json.dumps({"sql": s, "reason": reason}, ensure_ascii=False)


def _state(emp: str | None = EMP, q: str = "질문") -> dict:
    return {"question": q, "employee_id": emp}


def check_checker() -> None:
    ok = ["SELECT customer_name FROM v_irp_customer",
          "select count(*) from v_irp_customer where default_option = '미설정'",
          "SELECT TOP 3 customer_name FROM v_irp_customer ORDER BY eval_amt DESC;",
          "SELECT branch_name, target_count FROM v_branch_summary s",
          "SELECT customer_name FROM v_irp_customer WHERE target_reasons LIKE '%DELETE%'"]
    for s in ok:
        try:
            CK.check(s)
            passed = True
        except CK.SQLRejected as e:
            passed, s = False, f"{s} → {e}"
        check(passed, f"검사 통과: {s[:50]}", s)

    bad = ["DELETE FROM v_irp_customer",
           "SELECT * FROM irp_customer_base",
           "SELECT 1 FROM v_irp_customer; DROP TABLE branch",
           "SELECT customer_name FROM v_irp_customer -- 주석",
           "WITH v_irp_customer AS (SELECT * FROM irp_customer_base) SELECT * FROM v_irp_customer",
           "SELECT * INTO x FROM v_irp_customer",
           "SELECT * FROM v_irp_customer, irp_customer_base",
           "SELECT * FROM v_irp_customer WHERE age IN (SELECT age FROM employee)",
           "SELECT @@version FROM v_irp_customer",
           "SELECT * FROM sys.tables",
           "UPDATE v_irp_customer SET age = 1",
           ""]
    for s in bad:
        try:
            CK.check(s)
            rejected = False
        except CK.SQLRejected:
            rejected = True
        check(rejected, f"검사 거부: {s[:50] or '(빈 문자열)'}")


def check_scope() -> None:
    mine = run.my_branch(EMP)
    check(mine == ("0001", YEOUIDO), "사번 → 부점: 3902172 는 여의도종합금융센터", str(mine))
    check(run.my_branch("9999999") is None, "모르는 사번은 부점 없음")
    check(run.my_branch("12' OR '1'='1") is None, "사번 꼴이 아니면 조회하지 않음")

    _, rows, _ = run.execute("SELECT DISTINCT branch_name FROM v_irp_customer", "0001", EMP)
    check([r["branch_name"] for r in rows] == [YEOUIDO], "고객 뷰에는 직원 부점 고객만 있다",
          str(rows))
    # LLM 이 다른 부점을 적어도 범위 불변
    _, rows, _ = run.execute("SELECT customer_name FROM v_irp_customer WHERE branch_name = '강남'",
                             "0001", EMP)
    check(rows == [], "다른 부점을 조건으로 적어도 다른 부점 고객은 안 보인다", str(rows[:3]))
    _, rows, _ = run.execute("SELECT branch_name, is_my_branch FROM v_branch_summary", "0002",
                             OTHER_EMP)
    flags = {r["branch_name"]: r["is_my_branch"] for r in rows}
    check(len(rows) == 4 and flags.get("강남") == 1 and flags.get(YEOUIDO) == 0,
          "부점 비교는 집계 뷰로 — 내 부점 표시가 사번을 따른다", str(flags))


def check_limits() -> None:
    final, rows, cut = run.execute("SELECT TOP 3 customer_name FROM v_irp_customer ORDER BY eval_amt DESC",
                                   "0001", EMP)
    check(len(rows) == 3 and not cut and "LIMIT 3" in final, "TOP n → 목업 DB 에서는 LIMIT n",
          final[-80:])
    _, rows, cut = run.execute("SELECT customer_name FROM v_irp_customer", "0001", EMP)
    check(len(rows) <= run.MAX_ROWS, f"행 수 상한 {run.MAX_ROWS}")
    saved = run.MAX_ROWS
    try:
        run.MAX_ROWS = 10
        _, rows, cut = run.execute("SELECT customer_name FROM v_irp_customer", "0001", EMP)
        check(len(rows) == 10 and cut, "상한을 넘으면 자르고 잘렸다고 표시한다", f"{len(rows)} {cut}")
    finally:
        run.MAX_ROWS = saved
    check("TOP 101" in _sybase_limited("SELECT customer_name FROM v_irp_customer"),
          "Sybase 는 TOP 을 넣어 상한을 건다")
    check("TOP 5 " in _sybase_limited("SELECT TOP 5 customer_name FROM v_irp_customer"),
          "Sybase — 질문이 정한 TOP 은 그대로(상한 이하)")
    check("TOP 100 " in _sybase_limited("SELECT TOP 5000 customer_name FROM v_irp_customer"),
          "Sybase — 상한을 넘는 TOP 은 상한으로 줄인다")


def _sybase_limited(sql: str) -> str:
    saved = os.environ.get("BRANCH_DB")
    os.environ["BRANCH_DB"] = "sybase"
    try:
        return run._limited(sql)
    finally:
        if saved is None:
            os.environ.pop("BRANCH_DB", None)
        else:
            os.environ["BRANCH_DB"] = saved


def check_unavailable() -> None:
    def failed(state, env: str | None = None) -> str:
        saved = os.environ.get("BRANCH_DB")
        if env is not None:
            os.environ["BRANCH_DB"] = env
        try:
            with _LLM(_sql("SELECT customer_name FROM v_irp_customer")):
                tools.run("branch_search", state, "고객")
            return ""
        except ToolFailure as e:
            return e.reason
        finally:
            if saved is None:
                os.environ.pop("BRANCH_DB", None)
            else:
                os.environ["BRANCH_DB"] = saved

    from pension_agent.common.note import EMP_NO_ENV  # noqa: PLC0415
    saved_env = os.environ.pop(EMP_NO_ENV, None)
    try:
        check("부점" in failed(_state(None)), "사번이 없으면 조회하지 않는다(도구 고장)")
    finally:
        if saved_env is not None:
            os.environ[EMP_NO_ENV] = saved_env
    check("부점" in failed(_state("9999999")), "모르는 사번이면 조회하지 않는다")
    check("꺼져" in failed(_state(), "off"), "BRANCH_DB=off 면 조회하지 않는다")
    check("꺼져" in failed(_state(), "sqlite"), "표에 없는 값은 off 로 본다")
    check("작성되지" in failed(_state(), "sybase"), "Sybase 고정 뷰 SQL 이 TODO 면 조회하지 않는다")


def check_retry_and_evidence() -> None:
    with _LLM(_sql("SELECT customer_name FROM irp_customer_base"),
              _sql("SELECT customer_name, days_to_maturity FROM v_irp_customer "
                   "WHERE days_to_maturity BETWEEN 0 AND 30 ORDER BY days_to_maturity")) as llm:
        ev = tools.run("branch_search", _state(q="만기 30일 이내 고객 누구야?"), "만기 임박 고객")
    check(len(llm.prompts) == 2 and "고정 뷰만" in llm.prompts[1],
          "검사에 걸리면 이유를 실어 다시 쓰게 한다", str(len(llm.prompts)))
    check(ev is not None and ev["tool"] == "branch_search", "결과가 원장 근거가 된다")
    text = (ev or {}).get("text", "")
    check("| 고객명 | 만기까지 남은 일수 |" in text, "결과가 표로 실린다(머리글은 한글 이름)", text[:200])
    src = ((ev or {}).get("sources") or [{}])[0]
    check(src.get("id") == "branch_db" and "v_irp_customer" in (src.get("doc") or ""),
          "실행한 SQL 이 출처에 남는다", str(src))
    check(not re.search(r"\d{6}-\d{7}", text), "결과에 고객 식별번호가 없다")
    check("만기 30일 이내 고객 누구야?" in llm.prompts[0], "SQL 은 직원의 원문 질문으로 쓴다")
    check(views.catalog().splitlines()[0] in llm.prompts[0], "프롬프트에 고정 뷰 목록이 실린다")
    check("irp_customer_base" not in llm.prompts[0], "프롬프트에 원본 테이블 이름이 없다")

    with _LLM(_sql("SELECT * FROM branch")) as llm:
        try:
            tools.run("branch_search", _state(), "x")
            reason = ""
        except ToolFailure as e:
            reason = e.reason
    check(len(llm.prompts) == tool.RETRIES + 1 and "만들지 못했" in reason,
          f"끝내 검사에 걸리면 {tool.RETRIES + 1}번 시도 후 도구 고장", reason)

    # 조회할 수 없는 질문은 «확인한 결과»다 — 재료 0건으로 두면 계획이 다른 도구(관리 대상 명단)로 대신 답했다
    with _LLM(_sql(None, "날씨는 부점 데이터에 없음")) as llm:
        ev = tools.run("branch_search", _state(q="오늘 날씨 어때"), "날씨")
    text = (ev or {}).get("text", "")
    check(ev is not None and (ev.get("meta") or {}).get("cannot") and tool.CANNOT in text,
          "뷰로 답할 수 없다고 하면 «조회하지 못함»을 재료로 싣는다", text[:200])
    check("날씨는 부점 데이터에 없음" in text and tool.CANNOT_NO_SUBSTITUTE in text,
          "그 이유와 «다른 목록으로 대신하지 않는다»가 함께 실린다")
    check(len(llm.prompts) == 1, "답할 수 없다는 응답은 다시 쓰게 하지 않는다")
    with _LLM(_sql(None, "펀드 1천만원 이상 컬럼 없음")):
        ev = tools.run("branch_search", _state(q="펀드 1천만원 이상"), "펀드")
    check("1천만원" not in (ev or {}).get("text", ""),
          "숫자가 든 LLM 사유는 싣지 않는다(인용 허용을 넓히지 않는다)")
    with _LLM(_sql("ETF 보유 컬럼이 없어 조회할 수 없습니다")) as llm:
        ev = tools.run("branch_search", _state(q="x"), "x")
    check(ev is not None and (ev.get("meta") or {}).get("cannot") and len(llm.prompts) == 1,
          "sql 칸에 조회문 대신 설명을 적어도 «조회하지 못함»으로 읽는다(도구 고장 아님)")
    from pension_agent.prompts import ANSWER_SHAPES  # noqa: PLC0415
    check("조회하지 못함" in ANSWER_SHAPES["branch_search"] and "대신 답하지 않는다" in ANSWER_SHAPES["branch_search"],
          "답변 형태 요구가 «조회하지 못함»을 다른 목록으로 대신하지 말라고 한다")

    with _LLM(_sql("SELECT customer_name FROM v_irp_customer WHERE age > 200")):
        ev = tools.run("branch_search", _state(), "200세 이상")
    check(ev is not None and "0건" in ev["text"], "0건도 확인한 결과로 싣는다")

    t = tool.table([{"customer_name": "가", "return_rate": -9.4, "safe_asset_pct": None, "eval_amt": 5_000_000}])
    check("| -9.4% |" in t and "| 5,000,000 |" in t and "None" not in t and "|  |" in t,
          "뷰가 (%) 로 선언한 컬럼만 셀에 % 를 붙인다(답변의 「-9.4%」가 재료와 맞게)", t)
    with _LLM(_sql("SELECT customer_name, days_to_maturity, eval_amt FROM v_irp_customer "
                   "WHERE safe_asset_pct >= 80 AND days_to_maturity BETWEEN 0 AND 30")):
        crit_ev = tools.run("branch_search", _state(q="원리금보장 편중 · 만기 한 달 이내 고객"), "편중 만기 고객")
    check("· 조회 조건: 원리금보장·현금성 자산 비중 80% 이상 · 만기까지 남은 일수 0일~30일"
          in (crit_ev or {}).get("text", ""),
          "조회문의 조건이 재료에 말로 선다 — 백분율은 % 단위(답변의 «80% 이상»이 폐기되지 않게)",
          (crit_ev or {}).get("text", "")[:200])


def check_registry() -> None:
    check("branch_search" in tools.TOOLS, "branch_search 도구가 레지스트리에 있다")
    check("branch_search" in tools.usable({}), "고객 화면이 닫혀 있어도 쓸 수 있다")
    from pension_agent.prompts import ANSWER_SHAPES  # noqa: PLC0415
    check("branch_search" in ANSWER_SHAPES, "답변 형태 요구가 있다")


class _Today:
    """«오늘» 고정 + 목업 DB 재생성(합성 고객의 만기일이 오늘 기준이라 결과를 고정한다)."""

    def __init__(self, day: str = "2026-10-22"):
        self.day = day

    def __enter__(self):
        from pension_agent.common.clock import TODAY_ENV  # noqa: PLC0415
        from pension_agent.tools.branch_briefing import demo_db  # noqa: PLC0415
        self._env, self._saved = TODAY_ENV, os.environ.get(TODAY_ENV)
        os.environ[TODAY_ENV] = self.day
        demo_db.reset()
        return self

    def __exit__(self, *exc):
        from pension_agent.tools.branch_briefing import demo_db  # noqa: PLC0415
        if self._saved is None:
            os.environ.pop(self._env, None)
        else:
            os.environ[self._env] = self._saved
        demo_db.reset()


def check_rules() -> None:
    from pension_agent.tools.branch_briefing import status  # noqa: PLC0415
    r = status.rules()
    check(len(r.segments) >= 9 and len(r.cases) == 5, "부점 세그먼트·CASE 가 룰베이스 표에서 적재된다",
          f"{len(r.segments)}·{len(r.cases)}")
    check(all(s.conds for s in r.segments), "모든 세그먼트에 조건 행이 있다")
    check([c.priority for c in r.cases] == sorted(c.priority for c in r.cases), "CASE 는 우선순위 순이다")
    check("초안" in r.status, "표가 스스로 초안(검토 전)임을 밝힌다", r.status)
    bs01 = next(s for s in r.segments if s.id == "BS-01")
    check(status.segment_sql(bs01) == "idle_cash_pct >= 50", "BS-01 은 TG-001 과 같은 임계(고유계정대 50%)",
          status.segment_sql(bs01))
    check("납입금 미운용" in bs01.aliases and "현금성 미운용" in bs01.aliases,
          "「납입금 미운용」·「현금성 미운용」은 BS-01 의 다른 이름이다(용어 통일)")
    bad = [status.Segment("X", "x", "D", (("eval_amt; DROP", ">=", 1),), (), (), ""),
           status.Segment("X", "x", "D", (("eval_amt", "LIKE", 1),), (), (), ""),
           status.Segment("X", "x", "D", (("default_option", "=", "미설정' OR '1'='1"),), (), (), ""),
           status.Segment("X", "x", "D", (), (), (), "")]
    rejected = 0
    for seg in bad:
        try:
            status.segment_sql(seg)
        except status.RuleError:
            rejected += 1
    check(rejected == len(bad), "뷰에 없는 필드·연산자·따옴표 값·빈 조건은 SQL 로 옮기지 않는다",
          f"{rejected}/{len(bad)}")
    check("납입금 미운용" in tool.segment_catalog() and "WHERE idle_cash_pct >= 50" in tool.segment_catalog(),
          "고객 검색 SQL 프롬프트에 세그먼트 정의가 실린다")


def check_status() -> None:
    from pension_agent.tools.branch_briefing import status  # noqa: PLC0415
    with _Today():
        st = status.compute(EMP)
        other = status.compute(OTHER_EMP)
        ev = tools.run("branch_status", _state(), "부점 현황")
    ind = {i["id"]: i for i in st["indicators"]}
    check(st["branch"] == YEOUIDO and st["prev_date"] == "2026-09-30", "비교 기준은 전월말이다", st["prev_date"])
    check(ind["customer_count"]["direction"] == "증가" and ind["total_eval_amt"]["direction"] == "증가"
          and ind["branch_return_rate"]["label"] == "상승" and ind["idle_cash_total"]["direction"] == "감소",
          "지표 방향(화면 기획안 — 잔액 ▲ · 수익률 ▲ · 고유계정대 ▼)",
          str({k: v["direction"] for k, v in ind.items()}))
    check([c["id"] for c in st["cases"]] == ["BC-03"], "여의도종합금융센터는 CASE 3 이 성립한다",
          str(st["cases_all"]))
    check(other["cases"] == [], "강남은 성립하는 CASE 가 없다(0건도 결과)", str(other["cases_all"]))
    bc03 = st["cases"][0]
    check(bc03["segment_changes"][0]["id"] == "BS-06" and bc03["segment_changes"][0]["delta"] > 0,
          "CASE 는 판정에 쓴 세그먼트의 인원 변화를 싣는다", str(bc03["segment_changes"]))
    text = (ev or {}).get("text", "")
    check(ev is not None and "부점 현황" in text and "오늘의 관리 방향" in text, "branch_status 가 원장 근거가 된다")
    check("idle_cash_pct" not in text and "WHERE" not in text, "재료에 컬럼 이름·SQL 이 없다(개발 용어)")
    check(bool((ev or {}).get("marks")), "근거등급 D 기준이 섞이면 재료 성격 표시가 붙는다")
    check("branch_status" in tools.usable({}), "부점 현황은 고객 화면 없이 쓸 수 있다")

    # 임계값은 표가 정한다 — 세그먼트 증가 최소 인원을 올리면 CASE 가 사라진다
    r = status.rules()
    saved = dict(r.params)
    try:
        r.params["BP-01"] = 1000
        with _Today():
            st2 = status.compute(EMP)
    finally:
        r.params.clear()
        r.params.update(saved)
    check(st2["cases"] == [], "증가 판정 최소 인원(BP-01)을 올리면 CASE 가 성립하지 않는다")


def check_realistic() -> None:
    """여의도 고객군 인원이 현실적인 범위다 — 관리 대상 수에 맞추려고 나머지를 «깨끗하게» 만들지 않는다."""
    from pension_agent.tools.branch_briefing import status  # noqa: PLC0415
    with _Today():
        seg = {s["id"]: s["now"] for s in status.compute(EMP)["segments"]}
    for sid, lo, hi, what in (("BS-03", 10, 60, "현금성 장기대기"), ("BS-04", 40, 150, "만기도래 30일"),
                              ("BS-10", 150, 420, "디폴트옵션 미설정"), ("BS-08", 400, 800, "원리금보장 80% 이상"),
                              ("BS-09", 150, 450, "연금개시 가능 미개시")):
        check(lo <= seg[sid] <= hi, f"{what}({sid}) 인원이 {lo}~{hi}명 안이다", str(seg[sid]))


def check_etf() -> None:
    """ETF 평가금액 컬럼 — 목업 고객은 원장 보유상품의 ETF 합계, 합성 고객은 원리금보장 밖 몫 안."""
    from pension_agent.tools.customer_solution import customer as C  # noqa: PLC0415
    check(any(c == "etf_eval_amt" for c, _t, _d in views.columns("v_irp_customer")),
          "LLM 이 보는 뷰 목록에 ETF 평가금액이 있다")
    with _Today():
        _, rows, _ = run.execute("SELECT customer_name, etf_eval_amt, eval_amt, safe_asset_pct "
                                 "FROM v_irp_customer", "0001", EMP, limit=2000)
        by = {r["customer_name"]: r for r in rows}
        p = next(x for x in C.PERSONAS if x.nm == "최서윤")
        want = sum(h["amount"] for h in p.holdings if h.get("type") == "ETF")
        check(by["최서윤"]["etf_eval_amt"] == want, "목업 고객 ETF = 원장 보유상품 ETF 합계",
              f"{by['최서윤']['etf_eval_amt']} vs {want}")
        over = [r for r in rows
                if r["etf_eval_amt"] > r["eval_amt"] * (100 - (r["safe_asset_pct"] or 0)) / 100 + 10_000]
        check(not over, "ETF 는 원리금보장 밖 몫을 넘지 않는다", str(over[:2]))
        _, hit, _ = run.execute("SELECT customer_name FROM v_irp_customer "
                                "WHERE etf_eval_amt >= 10000000 AND return_rate < 5", "0001", EMP)
        names = {r["customer_name"] for r in hit}
        check(len(hit) >= 3 and "최서윤" in names,
              "«ETF 1천만원 이상 · 수익률 5% 미만»에 해당하는 고객이 있다", str(sorted(names)))
        _, prev, _ = run.execute("SELECT SUM(etf_eval_amt) AS s FROM v_irp_customer_prev", "0001", EMP)
        check((prev[0]["s"] or 0) > 0, "전월말 스냅샷에도 ETF 평가금액이 있다")


def check_maturity_type() -> None:
    """만기 상품 유형 컬럼 — 「정기예금 만기」를 아무 만기 상품으로 바꿔 조회하지 않게(2026-10-23 실측)."""
    from pension_agent.tools.branch_briefing import tool as BT  # noqa: PLC0415
    from pension_agent.tools.customer_solution import customer as C  # noqa: PLC0415
    check(any(c == "next_maturity_type" for c, _t, _d in views.columns("v_irp_customer")),
          "LLM 이 보는 뷰 목록에 만기 상품 유형이 있다")
    with _Today():
        _, rows, _ = run.execute("SELECT customer_name, next_maturity_type, next_maturity_date "
                                 "FROM v_irp_customer", "0001", EMP, limit=2000)
        by = {r["customer_name"]: r for r in rows}
        for nm in ("한지우", "김서연"):
            p = next(x for x in C.PERSONAS if x.nm == nm)
            first = sorted(p.maturities, key=lambda m: m.get("date") or "")[0]
            check(by[nm]["next_maturity_type"] == first["type"], f"목업 고객 {nm} — 만기 유형 = 원장 상품유형",
                  f"{by[nm]['next_maturity_type']} vs {first['type']}")
        odd = [r for r in rows if bool(r["next_maturity_date"]) != bool(r["next_maturity_type"])
               or r["next_maturity_type"] not in (None, "예금", "GIC")]
        check(not odd, "만기 상품이 있으면 유형이 '예금'·'GIC' 중 하나이고, 없으면 NULL", str(odd[:2]))
        kinds = {r["next_maturity_type"] for r in rows}
        check({"예금", "GIC"} <= kinds, "합성 고객에 정기예금과 GIC 가 함께 있다", str(kinds))
        _, prev, _ = run.execute("SELECT COUNT(*) AS n FROM v_irp_customer_prev WHERE next_maturity_type = '예금'",
                                 "0001", EMP)
        check(prev[0]["n"] > 0, "전월말 스냅샷에도 만기 유형이 있다")
    crit = BT.criteria("SELECT customer_name, days_to_maturity FROM v_irp_customer WHERE next_maturity_type = '예금' "
                       "AND days_to_maturity BETWEEN 0 AND 30 AND safe_asset_pct >= 80 AND default_option = '미설정' "
                       "ORDER BY days_to_maturity")
    check(crit == ["만기 상품: 정기예금", "만기까지 남은 일수 0일~30일", "원리금보장·현금성 자산 비중 80% 이상",
                   "디폴트옵션: 미설정"],
          "조회문의 조건을 전부 조회문 순서대로 말로 옮긴다 — 단위는 뷰 타입 선언에서(%·일)", str(crit))
    more = BT.criteria("SELECT TOP 5 customer_name FROM v_irp_customer WHERE etf_eval_amt >= 10000000 "
                       "AND target_reasons LIKE '%장기 미접촉%' AND is_target = 1")
    check(more == ["ETF 평가금액 1,000만원 이상", "관리 대상 요건 «장기 미접촉»(마지막 접촉 후 180일 이상)",
                   "관리 대상만"],
          "금액은 원 단위 표기 · 요건 사유는 그 요건의 기준과 함께 · 여부 컬럼은 «만»", str(more))
    cat = BT.segment_catalog()
    check("장기 미접촉(관리 대상 요건 — 마지막 접촉 후 180일 이상) → WHERE target_reasons LIKE '%장기 미접촉%'" in cat
          and "디폴트옵션 미설정(관리 대상 요건" not in cat,
          "SQL 프롬프트가 요건 이름의 정의를 싣는다 — «장기 미접촉»을 LLM 이 «90일 이상»으로 정하지 않게"
          "(세그먼트와 같은 이름의 요건은 싣지 않는다)")
    from pension_agent.tools.customer_solution import customer as CC  # noqa: PLC0415
    check(CC.DORMANT_DAYS == 180 and "180일" in cat, "장기 미접촉 기준은 판정 상수 하나에서 온다")
    t = BT.table([{"customer_name": "가", "days_to_maturity": -13}, {"customer_name": "나", "days_to_maturity": 5}])
    past = BT.criteria("SELECT a FROM v_irp_customer WHERE next_maturity_type = '예금' AND days_to_maturity < 0 "
                       "AND isa_days_to_maturity >= 0")
    check(past == ["만기 상품: 정기예금", "만기가 지남", "ISA 만기가 아직 안 지남"],
          "남은 일수 0 기준 조건은 «만기가 지남»으로 옮긴다(«0일 미만»이 아니라)", str(past))
    check("| -13 (만기 13일 지남) |" in t and "| 5 |" in t, "지난 만기는 칸에 «N일 지남»을 붙인다", t)
    check(BT.criteria("SELECT COUNT(*) AS n FROM v_irp_customer") == [], "조건이 없으면 조회 조건 줄도 없다")


def check_yeouido_screen() -> None:
    """부점 브리핑 화면 기획안의 수치가 목업 DB 에서 그대로 나온다(demo_yeouido)."""
    from pension_agent.tools.branch_briefing import status  # noqa: PLC0415
    with _Today():
        st = status.compute(EMP)
        _, rows, _ = run.execute("SELECT COUNT(*) AS n FROM v_irp_customer WHERE eval_amt = 100000000",
                                 "0001", EMP)
    ind = {i["id"]: i for i in st["indicators"]}
    want = {"customer_count": 1392, "total_eval_amt": 52_680_000_000, "branch_return_rate": 18.96,
            "idle_cash_total": 670_000_000, "safe_asset_ratio": 69.3, "transfer_out_count": 89,
            "transfer_out_amt": 6_280_000_000, "low_return_count": 210, "high_balance_count": 120,
            "high_balance_share": 45.3}
    got = {k: ind[k]["now"] for k in want}
    check(got == want, "화면 수치(1,392명 · 526.8억 · 18.96% · 6.7억 · 69.3% · 89좌·62.8억 · 210명 · 120명·45.3%)",
          str({k: v for k, v in got.items() if v != want[k]}))
    delta = (ind["total_eval_amt"]["delta"], ind["branch_return_rate"]["delta"], ind["idle_cash_total"]["delta"])
    check(delta == (130_000_000, 0.12, -40_000_000), "전월말 대비 ▲1.3억 · ▲0.12%p · ▼0.4억", str(delta))
    check(rows[0]["n"] == 0, "1억 경계에 몰린 계좌가 없다(비례 조정이 경계에 붙이지 않는다)")


def check_strategy() -> None:
    from pension_agent.tools.branch_briefing import status, strategy  # noqa: PLC0415
    with _Today():
        st = status.compute(EMP)
        none = status.compute(OTHER_EMP)
    nums = strategy.allowed_numbers(st)
    check(nums and all(n.endswith("명") for n in nums), "문장에 쓸 수 있는 숫자는 세그먼트 인원 변화뿐이다",
          str(nums))
    good = (f"부점 잔액이 늘었고 세액공제 한도 여유 고객도 {nums[0]} 늘었습니다. "
            "잔여한도가 크지만 최근 추가입금이 없는 고객부터 관리하는 것이 좋겠습니다.")
    check(strategy.check_text(good, st) == [], "방향만 쓴 지표 + 세그먼트 인원 변화는 통과한다")
    check(strategy.check_text("계좌 수가 6계좌 줄었습니다.", st) != [], "지표 수치를 쓰면 걸린다")

    saved = strategy.generate
    calls: list[str] = []

    def _gen(*replies):
        it = list(replies)

        def g(prompt, **kw):
            calls.append(prompt)
            return it.pop(0) if len(it) > 1 else it[0]
        return g
    try:
        strategy.generate = _gen(good)
        out = strategy.write(st)
        check(out["status"] == "ok" and out["text"] == good, "검사를 통과한 LLM 문장이 나간다")
        check("우선 관리할 고객" in calls[-1] and "잔여한도" in calls[-1], "프롬프트에 판정된 CASE 가 실린다")
        calls.clear()
        strategy.generate = _gen("잔액이 66억원으로 늘었습니다.", good)
        out = strategy.write(st)
        check(out["status"] == "ok" and len(calls) == 2 and "다음 이유" in calls[1],
              "검사에 걸리면 이유를 실어 한 번 다시 쓴다")
        strategy.generate = _gen("잔액이 66억원으로 늘었습니다.")
        out = strategy.write(st)
        check(out["status"] == "fallback" and out["text"] == " ".join(c["strategy"] for c in st["cases"]),
              "두 번 걸리면 기준표의 전략 문장을 그대로 싣는다")

        def down(prompt, **kw):
            raise strategy.LLMError("LLM 미설정 — 테스트")
        strategy.generate = down
        out = strategy.write(st)
        check(out["status"] == "llm_down" and out["text"] is None, "LLM 미연결이면 문장을 만들지 않는다(§11)")
        out = strategy.write(none)
        check(out["status"] == "no_case" and out["text"] == strategy.NO_CASE,
              "CASE 가 없으면 LLM 을 부르지 않고 코드 문장을 싣는다")
    finally:
        strategy.generate = saved


def check_targets() -> None:
    from pension_agent import graph  # noqa: PLC0415
    from pension_agent.tools.branch_briefing import roster, status  # noqa: PLC0415
    from pension_agent.tools.customer_solution import customer as C, target_list  # noqa: PLC0415
    from pension_agent.tools.customer_solution import tool  # noqa: PLC0415
    with _Today():
        _name, profiles = roster.branch_profiles(EMP)
        plain = target_list.today_targets(profiles)
        ev = tools.run("branch_targets", _state(q="오늘 중점적으로 관리할 고객은 누구야?"), "관리 대상")
        other = tools.run("branch_targets", _state(OTHER_EMP), "관리 대상")
        loose = tools.run("branch_targets", {"question": "q"}, "관리 대상")
        _, rows, _ = run.execute("SELECT customer_name FROM v_irp_customer WHERE is_target = 1 "
                                 "AND tax_credit_room >= 9000000 AND pension_started = '미개시' "
                                 "AND last_contact_days >= 180 AND eval_amt >= 30000000", "0001", EMP, limit=2000)
        pangyo = tools.run("branch_targets", _state("3902301"), "관리 대상")
        _, mats, _ = run.execute("SELECT customer_name FROM v_irp_customer WHERE is_target = 1 "
                                 "AND days_to_maturity BETWEEN 0 AND 30 AND next_maturity_amt >= 10000000 "
                                 "AND last_contact_days >= 90", "0003", "3902301", limit=2000)
    meta = (ev or {}).get("meta") or {}
    want = {r["customer_name"] for r in rows}
    check(len(plain) > 1000 and all(t.conds for t in plain),
          "현실적인 부점 — 관리 대상(요건 1개 이상)은 1,392명 중 1,000명을 넘는다", str(len(plain)))
    check(set(meta.get("names") or []) == want and 20 <= len(want) < 200 and YEOUIDO in ev["text"],
          "오늘의 중점관리 고객 = 관리 대상 ∩ CASE 3 중점관리 조건(AND)", f"{meta.get('count')} vs {len(want)}")
    check(meta.get("total") == meta.get("count") and f"관리 대상 {len(plain)}명" in ev["text"]
          and "900만원" in ev["text"], "재료가 관리 대상 전체 인원과 중점관리 조건을 밝힌다")
    check({n for n in pangyo["meta"]["names"]} == {r["customer_name"] for r in mats},
          "판교(CASE 1)는 CASE 1 의 중점관리 조건으로 고른다 — 부점 상황에 따라 다르다",
          f"{len(pangyo['meta']['names'])} vs {len(mats)}")
    persona_names = {p.nm for p in C.PERSONAS if C.conditions(p)}
    om = (other or {}).get("meta") or {}
    check("강남" in (other or {}).get("text", "") and not persona_names & set(om.get("names") or []),
          "다른 부점 직원에게는 그 부점 고객만 나온다(gap 39)")
    focus_max, plain_max = status.target_limits()
    check((focus_max, plain_max) == (100, 30), "인원 상한은 표(부점_임계값 BP-10 · BP-11)에서 온다 — 중점관리 100명",
          str((focus_max, plain_max)))
    check(om.get("count") == plain_max < om.get("total", 0) and "관리 방향이 없어" in other["text"],
          "관리 방향이 없는 부점(강남)은 관리 대상 전체를 우선순위 앞쪽부터 상한까지", str(om))
    # 상한을 넘으면 우선순위 앞쪽만 싣고 전체 인원을 밝힌다
    real_limits = status.target_limits
    status.target_limits = lambda: (10, plain_max)
    try:
        with _Today():
            capped = tools.run("branch_targets", _state(q="오늘 중점관리 고객"), "관리 대상")
    finally:
        status.target_limits = real_limits
    cm = capped["meta"]
    check(cm["count"] == 10 and cm["total"] == len(want) and cm["names"] == meta["names"][:10]
          and f"{len(want)}명 중 우선순위 앞쪽 10명만" in capped["text"],
          "중점관리 고객이 상한을 넘으면 우선순위 앞쪽만 싣고 전체 인원을 밝힌다", str(cm.get("count")))
    # CASE 는 섰는데 조건을 모두 갖춘 고객이 0명 — 중점관리가 아님을 밝히고 관리 대상 앞쪽을 따로 보인다
    real_focus = tool._today_focus
    tool._today_focus = lambda st, sc: (real_focus(st, sc)[0], set())
    try:
        with _Today():
            empty = tools.run("branch_targets", _state(q="오늘 중점관리 고객"), "관리 대상")
    finally:
        tool._today_focus = real_focus
    em, et = empty["meta"], empty["text"]
    check(em["count"] == plain_max and "중점관리 고객 0명" in et and "중점관리 고객이 아니라" in et
          and "(오늘의 중점관리 고객)" not in et.splitlines()[0] and "900만원" in et,
          "중점관리 0명이면 없다고 밝히고, 관리 대상 앞쪽을 중점관리와 구분해 보인다", et[:200])
    check(loose["meta"]["count"] == len(persona_names) and "·" in loose["text"].splitlines()[0],
          "사번이 없으면 예전처럼 목업 고객 명부로 답한다")
    check(ev is not None and meta.get("codes") == ["tax", "add"],
          "관리 대상 추천 순서는 오늘의 CASE 가 가리키는 요건을 따른다", str(meta.get("codes")))
    ordered = target_list.prioritize([t for t in plain if t.profile.nm in want], meta["codes"])
    check(meta["names"] == [t.profile.nm for t in ordered], "재료의 순서 = prioritize 결과")
    check(not re.search(r"\d{6}-\d{7}", ev["text"]) and "first_targets" in meta,
          "첫 추천은 이름으로만 남고 식별번호가 없다")

    first = graph._first_targets([], [ev])
    hist = [{"question": "오늘 누구 봐야 해", "first_targets": first}]
    check(graph._first_targets(hist, []) == first, "다음 턴은 이전 턴의 첫 추천을 잇는다")
    fake = dict(first, names=list(reversed(first["names"])), refs=list(reversed(first["refs"])))
    with _Today():
        again = tools.run("branch_targets", {"question": "처음 추천한 고객 다시 보여줘", "employee_id": EMP,
                                             "history": [{"question": "q", "first_targets": fake}]}, "다시")
    am = again["meta"]
    check(am["replayed"] and am["names"] == fake["names"] and "첫 추천" in again["text"],
          "이 대화의 첫 추천이 있으면 다시 계산하지 않고 그 목록을 그대로 싣는다")
    old = dict(first, date="2026-10-21")
    with _Today():
        fresh = tools.run("branch_targets", {"question": "q", "employee_id": EMP,
                                             "history": [{"question": "q", "first_targets": old}]}, "x")
    check(not fresh["meta"]["replayed"], "날짜가 바뀐 첫 추천은 다시 계산한다")
    from pension_agent.actions import memo  # noqa: PLC0415
    with _Today():
        table, _ = memo.table_for({"employee_id": EMP}, [again])
    # 표는 길이 상한에서 잘릴 수 있다(note.MAX_CHARS) — 앞의 두 사람 순서로 잰다
    a, b = table.find(fake["names"][0]), table.find(fake["names"][1])
    check(-1 < a < b, "쪽지 표도 도구가 실은 순서 그대로다")


def check_pension_savings() -> None:
    """연금저축 보유 컬럼 — 목업 고객은 원장 연금저축보유여부(비면 미확인 = NULL), 합성 고객은 시연 비율. 오늘·전월말 같은 값."""
    from pension_agent.tools.branch_briefing import demo_db  # noqa: PLC0415
    cols = [c for c, _t, _d in views.columns("v_irp_customer")]
    check("pension_savings" in cols, "연금저축 보유 컬럼이 고객 뷰에 있다")
    conn = demo_db.connection()
    vals = {r[0] for r in conn.execute("SELECT DISTINCT pension_savings FROM irp_customer_base")}
    check(vals <= {"보유", "미보유", None} and {"보유", "미보유"} <= vals, "연금저축 값은 보유·미보유·미확인", f"{vals}")
    diff = conn.execute("SELECT COUNT(*) FROM irp_customer_base a JOIN irp_customer_prev b "
                        "ON a.customer_name = b.customer_name AND a.branch_code = b.branch_code "
                        "WHERE a.pension_savings IS NOT b.pension_savings").fetchone()[0]
    check(diff == 0, "연금저축 보유는 오늘·전월말 표가 같다", f"{diff}명 다름")
    from pension_agent.tools.customer_solution import customer as C  # noqa: PLC0415
    known = {p.nm: {True: "보유", False: "미보유"}.get(p.pension_savings) for p in C.PERSONAS}
    got = dict(conn.execute("SELECT customer_name, pension_savings FROM irp_customer_base WHERE branch_code = '0001' "
                            f"AND customer_name IN ({','.join('?' * len(known))})", list(known)).fetchall())
    check(all(got.get(n) == v for n, v in known.items()), "목업 고객의 연금저축은 원장 값 그대로(비면 미확인)", f"{got}")


def check_demo_detail() -> None:
    """합성 고객 상세 원장(demo_detail) — 목록과 같은 값·같은 관리 사유 · 같은 고객이면 같은 값 · 목업 DB 일 때만."""
    from pension_agent.tools.branch_briefing import demo_db, demo_detail, roster  # noqa: PLC0415
    from pension_agent.tools.customer_solution import customer as C  # noqa: PLC0415
    names = {h["name"] for p in C.PERSONAS for h in p.holdings}
    keys = ("bal", "nm", "ag", "rk", "ret", "dopt", "room", "dorm", "pension_started", "pension_eligible",
            "club_grade", "cash_idle_pct", "port")
    for day in ("2026-10-22", "2026-10-23"):
        with _Today(day):
            rows = [dict(r, customer_ref=f"{r['branch_code']}:{r['rid']}") for r in
                    demo_db.connection().execute("SELECT rowid AS rid, * FROM irp_customer_base")]
            bad: list[str] = []
            n = 0
            for r in rows:
                ref = r["customer_ref"]
                if demo_db.detail_id(ref):
                    continue
                n += 1
                lst, det = roster.row_profile(r), C.get_profile(ref)
                if det is None:
                    bad.append(f"{ref} 상세 없음")
                    continue
                if C.conditions(lst) != C.conditions(det) or C.churn(lst) != C.churn(det):
                    bad.append(f"{ref} 요건 {C.conditions(lst)} ≠ {C.conditions(det)}")
                bad += [f"{ref} {k}" for k in keys if getattr(lst, k) != getattr(det, k)]
                if lst.matDate and (lst.matDate, lst.matAmt) != (det.matDate, det.matAmt):
                    bad.append(f"{ref} 만기")
                if (lst.isa or {}).get("amount") != (det.isa or {}).get("amount"):
                    bad.append(f"{ref} ISA")
                if sum(h["amount"] for h in det.holdings) != det.bal:
                    bad.append(f"{ref} 보유상품 합 ≠ 평가금액")
                if any(h["name"] not in names | {"고유계정대"} for h in det.holdings):
                    bad.append(f"{ref} 시연 고객 원장 밖 상품명")
                if any(h.get("discontinued") for h in det.holdings):
                    bad.append(f"{ref} 판매중단 상품")
            check(n > 1_500 and not bad, f"합성 고객 상세 원장이 목록 값·관리 사유와 같다({day} · {n}명)",
                  f"{len(bad)}건: {bad[:5]}")
            ref = next(r["customer_ref"] for r in rows if not demo_db.detail_id(r["customer_ref"]))
            first = C.get_profile(ref)
            demo_detail.clear()
            check(C.get_profile(ref) == first, "같은 고객이면 상세 원장이 같은 값이다(시드 = 내부 행 식별자)")
            persona_ref = next(r["customer_ref"] for r in rows if demo_db.detail_id(r["customer_ref"]))
            check(demo_detail.profile(persona_ref) is None and C.get_profile(demo_db.detail_id(persona_ref)),
                  "시연 고객은 자기 원장(customers.json)이다 — 생성기가 만들지 않는다")
            check(C.get_profile("0001:999999") is None and C.get_profile("모르는고객") is None,
                  "없는 행·모르는 id 는 상세가 없다")
            saved = os.environ.get("BRANCH_DB")
            os.environ["BRANCH_DB"] = "off"
            try:
                demo_detail.clear()
                check(C.get_profile(ref) is None, "목업 DB 가 아니면 합성 상세 원장을 만들지 않는다")
            finally:
                if saved is None:
                    os.environ.pop("BRANCH_DB", None)
                else:
                    os.environ["BRANCH_DB"] = saved
                demo_detail.clear()
    with _Today("2026-10-23"):
        from pension_agent.tools import customer_selection  # noqa: PLC0415
        ev = tools.run("branch_targets", _state(), "오늘의 중점관리 고객")
        listed = ev["meta"]["customer_list"]
        saved = {"employee_id": EMP, "session_id": "t", "serial": 1, "lists": [listed],
                 "active": listed["number"], "selected": None, "pending": None}
        out = customer_selection.resolve("1번 고객 자세히", saved, EMP, None)
        first = listed["entries"][0]
        check(out.customer_id == first["ref"] and out.case != "detail_unavailable",
              "목록에서 고른 합성 고객이 상세 원장으로 연결된다", f"{out.case} · {out.answer}")


def check_briefing_frame() -> None:
    """부점 브리핑 답변 틀(2026-10-23 사용자 결정) — 코드가 만든 제목·분석 기준·표 하나(부점 현황 | 값) 안에 LLM 문장.
    한 고객군만 물은 질문에는 틀을 붙이지 않는다. 찾은 고객 수는 굵게."""
    ev = tool._branch_status(_state(EMP, "오늘의 부점 브리핑 알려줘"), "오늘 부점 브리핑")
    frame = ev["meta"]["frame"]
    check(bool(frame), "부점 브리핑 질문에는 답변 틀이 붙는다")
    table = frame["table"]
    check(table.startswith("| 부점 현황 | 값 |"), "표의 열 이름은 «부점 현황 | 값»", table.splitlines()[0])
    for label in ("부점 잔액", "부점 수익률", "고유계정대 잔액", "원리금보장 비중", "올해 계약이전",
                  "우리부점 관리 고객 (개인형 IRP)", "수익률 저조 고객 (수익률 1% 미만)", "적립금 1억 이상 고객"):
        check(f"| {label} |" in table, f"표에 «{label}» 줄", table)
    check("(전체 적립금의 " in table.splitlines()[-1].split("|")[2], "적립금 비중은 값 열에 있다", table.splitlines()[-1])
    check(len([ln for ln in table.splitlines() if ln.startswith("|")]) == 10, "표는 하나(머리 2줄 + 8줄)")
    focus = ev["meta"]["focus"]
    if focus:
        check(f"오늘의 중점관리 고객: {focus}명" in ev["text"], "재료에 중점관리 고객 수")
        out = tool.framed(f"오늘은 고객 {focus}명을 찾았습니다.", frame)
        check(f"**고객 {focus}명**" in out, "찾은 고객 수는 굵게", out)
        check(out.startswith("> **오늘의 부점 브리핑** · ") and "분석" not in out.split("\n\n|")[0],
              "인용 상자: 제목 · 문장(분석 기준 줄 없음)")
        check(out.rstrip().endswith(table.splitlines()[-1]), "표는 인용 상자 아래")
    line = next((ln for ln in ev["text"].splitlines() if "올해 계약이전:" in ln), "")
    check("늘면 이탈이 늘었다는 뜻" in line, "재료가 계약이전 증가의 뜻(이탈)을 함께 싣는다", line)
    low = next((ln for ln in ev["text"].splitlines() if "수익률 저조 고객:" in ln), "")
    check("늘면 운용 관리가 필요한 고객이 늘었다는 뜻" in low, "재료가 수익률 저조 고객 증가의 뜻을 함께 싣는다", low)
    narrow = tool._branch_status(_state(EMP, "디폴트옵션 미설정 고객이 지난달보다 늘었어?"), "q")
    check(narrow["meta"]["frame"] is None, "한 고객군만 물은 질문에는 틀을 붙이지 않는다")
    check("화면 표" not in narrow["text"], "틀이 없으면 재료에 «화면 표» 줄도 없다")


def check_search_document() -> None:
    """고객 검색 답변 표 — 머리글이 DB 컬럼명이면 안 되고, 30명을 넘으면 엑셀로 전원을 붙인다.

    2026-10-23 실측: 멀티턴에서 «ETF 투자 금액 천만 원 이상 고객» 답이 «| 순서 | customer_name | age |
    etf_eval_amt | eval_amt |» 머리글로 73명 전부를 화면 표에 실었다."""
    import base64  # noqa: PLC0415
    import io  # noqa: PLC0415

    import openpyxl  # noqa: PLC0415

    from pension_agent.tools import customer_selection as CS  # noqa: PLC0415
    from pension_agent.tools.customer_solution import target_table  # noqa: PLC0415

    columns = {c for v in views.VIEWS for c, _t, _d in (views.VIEWS[v][1] or [])}

    def heads(doc: str) -> list[str]:
        return [c.strip() for line in doc.splitlines() if line.startswith("| 순서 |")
                for c in line.strip("|").split("|")]

    def ask(question: str, sql: str, saved: dict | None) -> tuple[dict, dict]:
        state = _state(q=question)
        if saved:
            state["customer_context"] = saved
        with _LLM(_sql(sql, "조회")):
            ev = tools.run("branch_search", state, question)
        return ev, CS.remember(saved or CS.context(None, EMP), [ev])

    with _Today():
        q1 = ("SELECT customer_ref, customer_name, age, etf_eval_amt, return_rate, safe_asset_pct, eval_amt "
              "FROM v_irp_customer WHERE etf_eval_amt >= 10000000 ORDER BY etf_eval_amt DESC")
        ev, saved = ask("ETF 투자 금액 천만 원 이상 고객 보여줘", q1, None)
        doc, meta = ev["meta"]["document"], ev["meta"]
        head = heads(doc)
        check(head and not set(head) & columns,
              "고객 검색 답변 표의 머리글은 DB 컬럼명이 아니라 한글 이름이다", f"{head}")
        check(head == ["순서", "고객명", "나이", "ETF 평가금액(원)", "운용 수익률", "원리금보장·현금성 자산 비중",
                       "평가금액(원)"], "금액 열은 «(원)»을 붙이고 백분율 열은 셀에 % 가 붙는다", f"{head}")
        check("조회 조건: ETF 평가금액 1,000만원 이상" in doc, "답변에 조회 조건이 말로 선다", doc[:300])
        n = meta["rows"]
        rows = [ln for ln in doc.splitlines() if ln.startswith("| ")][1:]
        check(len(rows) == min(n, target_table.TEXT_ROWS),
              f"화면 표는 {target_table.TEXT_ROWS}명까지", f"{len(rows)} / {n}")
        files = meta.get("files") or []
        check(n > target_table.TEXT_ROWS and len(files) == 1 and files[0]["name"] in doc
              and f"{n}명 전체는 첨부한 엑셀" in doc,
              "30명을 넘으면 전원을 엑셀로 붙이고 본문이 파일 이름을 밝힌다", f"{n} · {files and files[0]['name']}")
        if files:
            sheet = list(openpyxl.load_workbook(io.BytesIO(base64.b64decode(files[0]["base64"]))).active
                         .iter_rows(values_only=True))
            check(list(sheet[0]) == head and len(sheet) - 1 == n
                  and [r[0] for r in sheet[1:]] == list(range(1, n + 1))
                  and all(isinstance(r[3], int) for r in sheet[1:]),
                  "엑셀은 같은 머리글·같은 순서 번호로 전원이고 금액은 숫자 칸이다", f"{sheet[0]} · {len(sheet) - 1}")
            check(not any("customer_ref" in str(c) or re.fullmatch(r"\d{4}:\d+", str(c)) for r in sheet for c in r),
                  "엑셀에도 내부 행 식별자가 없다")

        # 엑셀에만 있는 번호로 이어서 고른다 — 목록은 표에 자른 30명이 아니라 전원을 기억한다
        pick = CS.resolve("35번 고객 자세히", saved, EMP)
        entry = saved["lists"][-1]["entries"][34]
        check(pick.customer_id == entry["ref"], "엑셀의 35번을 고르면 그 고객이 선택된다",
              f"{pick.case} · {pick.customer_id} vs {entry['ref']}")

        # 이어지는 턴(멀티턴) — 좁힌 조회도 같은 꼴이고 30명 이하면 첨부가 없다
        q2 = ("SELECT customer_ref, customer_name, etf_eval_amt, safe_asset_pct FROM v_irp_customer "
              "WHERE etf_eval_amt >= 10000000 AND safe_asset_pct >= 80")
        ev2, saved = ask("그 중에 원리금보장 비중 80% 이상인 고객은?", q2, saved)
        doc2 = ev2["meta"]["document"]
        check(not set(heads(doc2)) & columns and ev2["meta"]["customer_list"]["number"] == 2,
              "이어지는 턴의 목록도 머리글이 한글이고 목록 번호가 이어진다", f"{heads(doc2)}")
        check(0 < ev2["meta"]["rows"] <= target_table.TEXT_ROWS and not ev2["meta"].get("files"),
              "30명 이하 목록에는 엑셀을 붙이지 않는다", str(ev2["meta"]["rows"]))
        ev3, _ = ask("ETF 천만 원 이상이면서 원리금보장 90% 이상 고객", q2.replace(">= 80", ">= 90"), saved)
        check("(조회 결과 0건)" in ev3["text"] and not ev3["meta"].get("files"),
              "0건도 0건이라고 싣고 첨부가 없다")


def check_question_set() -> None:
    """시험 질문(tests/branch_questions.py) — LLM 없이 확인할 수 있는 것 전부.
    정답 조회문이 통과·실행되는지 · 행 수 · «조회 조건» 줄 · SQL 프롬프트의 기준 · 이상적인 답이 검증기를 통과하는지.
    라이브(실제 LLM) 채점은 python -m tests.live_branch."""
    from pension_agent.evidence.verify import verify_texts  # noqa: PLC0415
    from pension_agent.common import clock  # noqa: PLC0415
    from tests import branch_questions as BQ  # noqa: PLC0415
    prompt = views.catalog() + "\n" + tool.segment_catalog()
    ids = [q.id for q in BQ.QUESTIONS]
    check(len(ids) == len(set(ids)), "시험 질문: id 중복 없음")
    for q in BQ.QUESTIONS:
        for text in q.prompt_has:
            check(text in prompt, f"시험 질문 {q.id}: SQL 프롬프트에 «{text}»")
        if "branch_search" not in q.tool.split("|"):
            continue
        with _LLM(BQ.llm_reply(q, "조회" if q.sql else "그 항목의 컬럼이 없습니다")) as llm:
            ev = tool._branch(_state(BQ.EMP, q.question), q.question)
        today = f"<오늘>\n{clock.today().isoformat()}\n</오늘>"
        check(all(today in p for p in llm.prompts), f"시험 질문 {q.id}: SQL 프롬프트에 오늘 날짜")
        text = ev["text"] if ev else ""
        if q.sql is None:
            check(ev and ev["meta"].get("cannot"), f"시험 질문 {q.id}: 조회할 수 없음 재료")
        else:
            got = tool.criteria(q.sql)
            check(got == list(q.criteria), f"시험 질문 {q.id}: 조회 조건 줄", f"{got}")
            check(ev and ev["meta"]["rows"] == q.rows and ev["meta"]["truncated"] == q.truncated,
                  f"시험 질문 {q.id}: 행 수 {q.rows}",
                  f"{ev and ev['meta']['rows']} · 잘림 {ev and ev['meta']['truncated']}")
        if q.ideal:
            ok, bad = verify_texts(q.ideal, ev["allow"] if ev else [], echoable=[q.question])
            check(ok, f"시험 질문 {q.id}: 이상적인 답이 검증 통과", f"{bad}")
            faults = BQ.answer_faults(q, q.ideal)
            check(not faults, f"시험 질문 {q.id}: 이상적인 답이 채점 기준 통과", f"{faults}")
        if q.sql:
            check("```sql" in text, f"시험 질문 {q.id}: 재료에 실행한 조회문")


def check_target_table() -> None:
    from dataclasses import replace  # noqa: PLC0415
    from unittest.mock import patch  # noqa: PLC0415

    from pension_agent.common import note  # noqa: PLC0415
    from pension_agent.nodes import answer as ANS, plan  # noqa: PLC0415
    from pension_agent.tools.customer_solution import customer as C, target_table  # noqa: PLC0415
    from pension_agent.tools.customer_solution.target_list import Target  # noqa: PLC0415

    with _Today():
        ev = tools.run("branch_targets", _state(), "관리 대상")
    doc = ev["meta"]["document"]
    rows = [line for line in doc.splitlines() if line.startswith("| ")][1:]
    shown = min(ev["meta"]["count"], target_table.TEXT_ROWS)
    check(len(rows) == shown and doc.count("| 순서 |") == 1,
          f"표 하나에 앞쪽 {target_table.TEXT_ROWS}명까지 나온다(카드·분할 없음)", f"{len(rows)} vs {shown}")
    files = ev["meta"]["files"]
    check(len(files) == 1 and files[0]["name"].endswith(".xlsx") and files[0]["name"] in doc
          and files[0]["mime"] == target_table.XLSX_MIME,
          "관리 대상 목록에는 엑셀 파일이 하나 붙고 본문이 파일 이름을 밝힌다")
    if ev["meta"]["count"] > target_table.TEXT_ROWS:
        check(f"앞쪽 {target_table.TEXT_ROWS}명만" in doc and f"{ev['meta']['count']}명 전체는 첨부한 엑셀" in doc,
              "표에서 자른 인원은 엑셀에 있다고 밝힌다")
    import base64, io  # noqa: E401, PLC0415
    import openpyxl  # noqa: PLC0415
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(files[0]["base64"]))).active
    sheet = list(ws.iter_rows(values_only=True))
    check(list(sheet[0]) == target_table.SHEET_HEADER and len(sheet) - 1 == ev["meta"]["count"]
          and [r[1] for r in sheet[1:]] == ev["meta"]["names"] and [r[0] for r in sheet[1:]] == list(
              range(1, ev["meta"]["count"] + 1)),
          "엑셀은 실은 전원을 같은 순번·순서로 싣는다", f"{len(sheet) - 1} vs {ev['meta']['count']}")
    check(all(isinstance(r[6], int) and isinstance(r[8], int) and isinstance(r[7], (int, float)) for r in sheet[1:])
          and ws.freeze_panes == "A2" and ws.auto_filter.ref,
          "엑셀 금액·수익률은 숫자 셀이고 머리글 고정·필터가 있다")
    check(not any(re.search(r"\d{6}-\d{7}", str(v)) for r in sheet for v in r),
          "엑셀에도 PIN 은 마스킹만 싣는다")
    m = re.search(r"모든 고객에게 같은 사유\(([^)]*(?:\([^)]*\)[^)]*)*)\)는 표에서", doc)
    dropped = [x.strip() for x in m.group(1).split(", ")] if m else []
    check(bool(dropped) and all(all(x in r[12] for x in dropped) for r in sheet[1:]),
          "엑셀 관리 사유는 표에서 뺀 공통 사유도 싣는다(행 하나만 봐도 사유가 다 보이게)", f"{dropped}")
    check(all(f"| {i} |" in row and name in row
              for i, (name, row) in enumerate(zip(ev["meta"]["names"], rows), 1)),
          "표의 고객·순번은 추천 순서를 그대로 보존한다")
    check(not re.search(r"\d{6}-\d{7}", doc) and "고객 PIN" in doc and "미확인" in doc,
          "PIN 은 마스킹하고 PIN 없는 합성 고객의 내부 키는 미확인이다")
    check("공제한도 소진율" in doc and "공제 잔여한도(소진율)" in doc and "IRP 잔액(1년 수익률)" in doc
          and "900만원" in doc,
          "소진율·잔여한도와 기본한도 기준이 함께 표시된다")

    p = C.PERSONAS[0]
    samples = [Target(replace(p, room=room, ret=ret, pension_savings=True), ["mat", "hlt", "dep", "nod"])
               for room, ret in ((0, 6.2), (900, 2.8), (300, 21.6), (550, -9.4))]
    rendered = target_table.render(samples)
    sample_rows = rendered.splitlines()[2:]
    check(all(f"(소진 {pct}%) |" in row for row, pct in zip(sample_rows, (100, 0, 67, 39)))
          and "(-9.4%) |" in sample_rows[-1] and "(+6.2%) |" in sample_rows[0],
          "소진율 경계·반올림과 수익률의 양수·음수 부호가 정확하다")
    alone = [target_table.render([t]).splitlines()[2] for t in samples]
    check(all(all(r in row for r in target_table.reasons(t)) and "연금저축 보유" in row
              for t, row in zip(samples, alone)),
          "관리 사유는 세 건을 넘어도 생략하지 않고 연금저축 보유도 유지한다")
    check(all("편중(80% 이상)" not in r and "만기예금 보유" not in r for r in alone),
          "관리 사유는 짧은 이름으로 쓴다(«원리금보장 N%» · «만기예금 D-n(금액)»)")
    same = target_table.common(samples)
    check(same and all(not any(f", {x}" in row or f"| {x}" in row for x in same) for row in sample_rows),
          "모든 고객에게 글자까지 같은 사유는 표에서 뺀다", f"{same}")
    check("모든 고객에게 같은 사유" in doc, "뺀 공통 사유는 표 위에서 한 번 밝힌다")
    merged = Target(replace(p, rk="위험중립형", port=[20, 6, 70, 4], pension_savings=False),  # 위험자산 74%
                    ["lim", "mis", "nod"])
    check(target_table.reasons(merged) == ["위험자산 74%(한도 초과 · 성향 불일치)", "디폴트옵션 미설정"],
          "같은 비중을 두 번 말하지 않는다(위험자산 한도 초과 + 성향 불일치 → 한 칸)",
          str(target_table.reasons(merged)))
    dirty = Target(replace(p, nm="고객|이름\n<태그>", id="0001:내부키", club_grade=None), ["tax"])
    dirty_row = target_table.render([dirty]).splitlines()[2]
    check(len(re.split(r"(?<!\\)\|", dirty_row)) == 9 and "<태그>" not in dirty_row
          and "0001:내부키" not in dirty_row and "미확인" in dirty_row,
          "셀의 구분자·줄바꿈·HTML 과 내부 키가 표를 깨뜨리거나 PIN 이 되지 않는다")
    check(target_table.render([]) == note.EMPTY_BODY, "0명은 빈 표 대신 대상 없음으로 안내한다")

    with patch.object(plan, "generate", return_value=_sql(None)) as unused:
        with patch.object(ANS, "compose", side_effect=AssertionError("표를 재작성하면 안 됨")):
            out = ANS.answer({"question": "오늘 누구부터 봐야 해?", "evidence": [ev]})
    check(out["answer"] == doc and not unused.called,
          "완성된 표는 모델 재작성 없이 그대로 답변이 된다")
    check(not re.search(r"\d{6}-\d{7}", ev["text"] + str(ev["sources"]) + str(ev["allow"]))
          and "고객 PIN" not in ev["text"],
          "모델·관측용 근거와 출처는 PIN 표시 문서와 분리된다")

    with _Today(), patch.object(plan, "generate", return_value='{"tool":"branch_targets","query":"관리 대상"}'):
        result = plan.plan_step({"question": "오늘 누구부터 봐야 해?", "employee_id": EMP})
    check(result.get("plan_done") and result["evidence"][0]["meta"].get("document"),
          "관리 대상 표가 완성되면 재료 수집을 종료해 누락·재요약을 막는다")


def main() -> int:
    check_checker()
    check_scope()
    check_limits()
    check_unavailable()
    check_retry_and_evidence()
    check_registry()
    check_rules()
    check_status()
    check_yeouido_screen()
    check_etf()
    check_maturity_type()
    check_realistic()
    check_strategy()
    check_targets()
    check_pension_savings()
    check_demo_detail()
    check_briefing_frame()
    check_target_table()
    check_search_document()
    check_question_set()

    passed = sum(1 for ok, _, _ in _results if ok)
    total = len(_results)
    print(f"\n{passed}/{total} 통과")
    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
