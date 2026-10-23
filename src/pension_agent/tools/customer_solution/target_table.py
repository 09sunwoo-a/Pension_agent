"""관리 대상 목록의 마크다운 표 — 선정·정렬 없이 원장 값만 표시."""

from __future__ import annotations

import html
import re

from pension_agent.common import note
from pension_agent.tools.customer_solution.customer import TAX_CREDIT_CAP_WON, Profile
from pension_agent.tools.customer_solution.target_list import Target


def _cell(value: str) -> str:
    """셀 안의 줄바꿈·마크다운 구분자·HTML 이 표 구조를 바꾸지 않도록 처리."""
    text = html.escape(" ".join(value.split()), quote=False)
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("`", "\\`").replace("*", "\\*")


def _pin(profile: Profile) -> str:
    """원장 PIN 만 표시. 부점 DB 의 내부 키(부점:이름)는 PIN 으로 쓰지 않음."""
    if not re.fullmatch(r"\d{6}-\d{7}", profile.id):
        return "미확인"
    return note.display_customer_id(profile.id)


#: 관리 사유 칸의 짧은 요건 이름(2026-10-23 사용자 결정 — «원리금보장상품 편중(80% 이상) 100%» → «원리금보장 100%»).
#: 값은 note.DETAIL 그대로이고 이름만 줄인다. 여기 없는 요건은 CONDS 이름 그대로.
SHORT_CONDS = {
    "chn": "이탈 위험 높음", "nod": "디폴트옵션 미설정", "pen": "연금개시 요건충족 후 미개시",
    "nch": "운용변경 없음", "sec": "섹터 집중", "hlt": "판매중단·환매추천 펀드", "idl": "미운용 현금",
    "isa": "ISA 만기자금", "out": "이탈위험 관찰", "tax": "세액공제 여력", "add": "추가입금 여력",
    "low": "수익률 하위 30%", "dor": "미접촉", "mat": "만기예금",
}
SAVINGS = "연금저축 보유"


def _paren(value: str) -> str:
    """«D-7 · 1,926만원» → «D-7(1,926만원)»."""
    head, sep, rest = value.partition(" · ")
    return f"{head}({rest})" if sep else value


def reasons(target: Target) -> list[str]:
    """관리 사유 — 짧은 이름 · 같은 값을 두 번 말하지 않는다(투자성향 불일치가 보여주는 비중이 원리금보장·위험자산
    한도와 같은 값이면 그 값 옆에 «성향 불일치»로 붙인다). 순서는 요건 순서 그대로."""
    p = target.profile
    conds = list(target.conds)
    metric = note.DETAIL["mis"](p).partition(" · ")[2] if "mis" in conds else ""
    shown = {"dep": f"원리금보장 {p.port[0]}%", "lim": f"위험자산 {p.risk_asset}%"}
    out = []
    for c in conds:
        if c in shown:
            tags = (["한도 초과"] if c == "lim" else []) + (["성향 불일치"] if metric == shown[c] else [])
            out.append(f"{shown[c]}({' · '.join(tags)})" if tags else shown[c])
        elif c == "mis":
            if not any(k in conds and shown[k] == metric for k in shown):
                out.append(f"투자성향 불일치({metric})")
        else:
            try:
                value = note.DETAIL.get(c, lambda _p: None)(p)
            except Exception:  # noqa: BLE001
                value = None
            name = SHORT_CONDS.get(c, note.CONDS.get(c, c))
            out.append(f"{name} {_paren(value)}" if value else name)
    if p.pension_savings:
        out.append(SAVINGS)
    return out


def common(targets: list[Target]) -> list[str]:
    """모든 고객에게 글자까지 같은 사유 — 표에서 빼고 표 위에서 한 번 말한다(사람마다 값이 다르면 남는다)."""
    if len(targets) < 2:
        return []
    each = [reasons(t) for t in targets]
    return [r for r in each[0] if all(r in rs for rs in each[1:])]


def render(targets: list[Target]) -> str:
    """선정 순서를 유지한 표 하나. PIN 은 기존 쪽지의 마스킹 규칙 적용.
    잔액과 1년 수익률 · 공제 잔여한도와 소진율은 한 칸에 괄호로 둔다(2026-10-23 사용자 결정 — 열 수를 줄인다).
    관리 사유는 짧은 이름으로 쓰고, 모든 고객에게 같은 사유는 뺀다(common — 호출하는 쪽이 표 위에 밝힌다)."""
    if not targets:
        return note.EMPTY_BODY
    same = set(common(targets))
    lines = [
        "| 순서 | 고객 PIN | 고객(나이·등급) | 투자성향 | IRP 잔액(1년 수익률) | 공제 잔여한도(소진율) | 관리 사유 |",
        "|---:|---|---|---|---:|---:|---|",
    ]
    for i, target in enumerate(targets, 1):
        p = target.profile
        customer = " · ".join([p.nm, f"{p.ag}세"] + ([p.club_grade] if p.club_grade else []))
        used = (1 - p.room * 10_000 / TAX_CREDIT_CAP_WON) * 100
        rest = [r for r in reasons(target) if r not in same]
        cells = [str(i), _pin(p), customer, p.rk, f"{note.won(p.bal)} ({p.ret:+.1f}%)",
                 f"{note.won(p.room * 10_000)} (소진 {used:.0f}%)",
                 ", ".join(rest) or "—"]
        lines.append("| " + " | ".join(_cell(value) for value in cells) + " |")
    return "\n".join(lines)


#: 엑셀 첨부의 열(2026-10-23 사용자 요청 — 행내 «타겟고객 명단» 엑셀 꼴). 표와 달리 값마다 한 칸이고 숫자는 숫자 셀이다
#: (정렬·필터·합계를 엑셀에서 하게). 관리 사유는 공통 사유도 빼지 않는다 — 행 하나만 떼어 봐도 사유가 다 보이게.
#: 성별은 원장에 없어 열을 두지 않는다(지어내지 않는다).
SHEET_HEADER = ["순번", "고객명", "고객 PIN", "나이", "스타클럽 등급", "투자성향", "IRP 잔액(원)", "1년 수익률(%)",
                "세액공제 잔여한도(원)", "공제한도 소진율(%)", "디폴트옵션", "원리금보장 비중(%)", "관리 사유"]


def sheet_rows(targets: list[Target]) -> list[list[object]]:
    """엑셀 첨부의 행 — 선정 순서 그대로. 값은 원장 값(PIN 은 표와 같은 마스킹)."""
    rows = []
    for i, target in enumerate(targets, 1):
        p = target.profile
        used = round((1 - p.room * 10_000 / TAX_CREDIT_CAP_WON) * 100, 1)
        rows.append([i, p.nm, _pin(p), p.ag, p.club_grade or "", p.rk, int(p.bal), round(float(p.ret), 1),
                     int(p.room * 10_000), used, p.dopt, int(p.port[0]), ", ".join(reasons(target))])
    return rows


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: 텍스트 표에 싣는 최대 인원(2026-10-23 사용자 요청 — 나머지는 엑셀 첨부). 답변 형태 규칙이라 코드 상수다.
TEXT_ROWS = 30


def workbook(targets: list[Target], title: str) -> bytes:
    """관리 대상 목록 → .xlsx 바이트(시트 하나). 머리글 굵게·회색 · 머리글 고정 · 자동 필터 · 숫자 서식."""
    import io  # noqa: PLC0415

    from openpyxl import Workbook  # noqa: PLC0415 — 이 기능에서만 쓴다
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(SHEET_HEADER)
    for row in sheet_rows(targets):
        ws.append(row)
    line = Side(style="thin", color="BFBFBF")
    border = Border(left=line, right=line, top=line, bottom=line)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.border = border
            if isinstance(cell.value, float):
                cell.number_format = "0.0"
            elif isinstance(cell.value, int) and cell.row > 1:
                cell.number_format = "#,##0"
    for i, head in enumerate(SHEET_HEADER, 1):
        values = [head] + [ws.cell(r, i).value for r in range(2, ws.max_row + 1)]
        width = max(sum(2 if ord(ch) > 0x2E80 else 1 for ch in (f"{v:,}" if isinstance(v, int) else str(v or "")))
                    for v in values)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 80)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
