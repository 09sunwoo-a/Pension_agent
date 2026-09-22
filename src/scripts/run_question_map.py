# -*- coding: utf-8 -*-
"""질문 리스트를 실 LLM 으로 돌려 xlsx 의 「실측 답변」 칸을 채운다.

    cd src
    python -m scripts.run_question_map --new --grade ◎        # 기존 대본에 없던 1순위 전부
    python -m scripts.run_question_map --rows 17,54,62         # 번호로 골라서
    python -m scripts.run_question_map --rows 10-20            # 구간 (양끝 포함 · 10~20 · 10–20 도 같다)
    python -m scripts.run_question_map --item 세액공제 --full   # 항목으로 · 근거까지 기록
    python -m scripts.run_question_map --new --dry-run         # 무엇을 돌릴지만 본다
    python -m scripts.run_question_map                         # 아직 안 채운 행 전부
    python -m scripts.run_question_map --retry-failed          # 실패로 끝난 칸만 다시
    python -m scripts.run_question_map --auto                  # 답이 대조로 갈리는 항목만
    python -m scripts.run_question_map --manual                # 사람이 읽어야 하는 항목만

**기록하는 것은 답변 본문뿐이다.** 진행 표시(「제도·상품 수치 찾는 중…」)와 근거 목록·
화면 딥링크는 화면에만 찍고 칸에는 안 넣는다 — 셀 하나가 근거 덤프로 길어지면 표를
읽을 수 없게 되고, 근거는 어차피 트레이스(`python -m tests.debug --debug`)에서 본다.
함께 남기고 싶으면 `--full`(근거+화면링크) · `--sources`(근거만) · `--links`(화면만).

읽고 쓰는 파일은 `docs/QUESTION_MAP.xlsx` 하나다. 질문·대상 고객·선행 질문을 그 표에서
읽고, 같은 행의 「실측 답변」 칸에 답을 써서 다시 저장한다. **질문 자체를 고치려면 여기가
아니라 `scripts/question_map.py` 다**(xlsx 는 산출물이다 — 루트 CLAUDE.md 규칙 3).
생성기는 실측 답변을 보존하므로 이 순서로 섞어 써도 된다.

    python -m scripts.run_question_map ...     # 답을 채우고
    python -m scripts.question_map             # 질문을 고쳐 다시 만들어도 답은 남는다

━━ 돌리기 전에 ━━
① **오늘을 고정한다.** 「대상 고객」 칸이 2026-09-29 기준 계산값이라, 다른 날로 돌리면
   만기·연금개시 요건이 달라져 기대와 어긋난다. 이 스크립트가 `PENSION_TODAY` 를 안 잡아
   주지는 않는다 — 환경변수로 준다(기본값을 코드에 박으면 실데이터 전환 때 조용히 틀린
   날짜로 돈다).

       PENSION_TODAY=2026-09-29 python -m scripts.run_question_map --new --grade ◎

② **브리핑을 미리 만든다.** 고객을 여는 첫 턴은 브리핑 생성(LLM 11회)을 치른다. 분당
   한도가 있는 키면 그 11회가 한도를 소진해 답변 작성이 429 로 죽는다.
   `python -m scripts.prebuild_briefings` 를 한 번 돌려 두면 읽어 쓴다. 그래도 나면
   `--pause=20`.

③ **상담 기록을 더럽히지 않는다.** 실행은 `tests.debug.runner.session()` 을 그대로 쓰므로
   세션마다 기록이 남았다가 나갈 때 걷힌다(그 함수 주석). `git status src/session_data/`
   가 깨끗한지 끝나고 한 번 본다.

④ **쪽지 행(「이번 상담 요약·쪽지」)은 사번이 있어야 돈다.** 실서비스에서는 프론트가
   로그인 사번을 넘기지만 이 실행기는 안 넘긴다 — 그러면 에이전트는 받을 사람을 모르고,
   기준서 §10 「받을 사람을 모르면 묻지 않는다」대로 **제안 자체를 하지 않는다.** 그게
   올바른 동작이라 화면에는 오류가 안 뜨고, 칸만 「쪽지를 보낼 받는 사람을 알 수 없어요」로
   채워진다(108·109 가 그렇게 채워져 있었다). 폴백 환경변수로 준다.

       WORKB_EMP_NO=3902172 PENSION_TODAY=2026-09-29 python -m scripts.run_question_map --rows 108,109

   **발송까지 가지는 않는다** — `MCP_*` 가 없으면 WorkB 클라이언트가 안 붙어 「미연결」로
   답한다(109 의 확인 포인트가 재는 것이 그 자리다).
"""

from __future__ import annotations

import argparse
import os
import sys
import time

from pension_agent import config
from scripts import question_map as QM

XLSX = QM.XLSX

#: 표에서 읽는 열. 없으면 그 자리에서 멈춘다 — 열 이름이 바뀐 것을 조용히 넘기면
#: 엉뚱한 칸에 답을 쓴다.
COLS = ("항목", "번호", "질문", "실측 답변", "시연", "대상 고객", "선행 질문")

#: 실행이 죽은 칸의 머리말. 이 글로 시작하는 칸은 «채워진 것»으로 세지 않는다 — 그래야
#: 다음 실행에서 다시 집히고(`--retry-failed`), 성공한 답을 덮어쓰지 않는다.
FAILED = "[실행 실패]"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def _load():
    """(워크북, 시트, 열이름→인덱스(1부터))."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        raise SystemExit(
            "openpyxl 이 필요하다: pip install openpyxl\n"
            "(배포 이미지에는 안 들어간다 — 이 스크립트는 개발용이다)") from None
    if not XLSX.is_file():
        raise SystemExit(f"{XLSX} 가 없다. 먼저: python -m scripts.question_map")
    wb = load_workbook(XLSX)
    ws = wb.active
    head = [c.value for c in next(ws.rows)]
    idx = {}
    for name in COLS:
        if name not in head:
            raise SystemExit(f"xlsx 에 「{name}」 열이 없다 — 표를 다시 만든다: "
                             "python -m scripts.question_map")
        idx[name] = head.index(name) + 1
    return wb, ws, idx


def _filled(ws, col: int) -> dict[int, str]:
    """세로 병합된 칸을 앞 값으로 이어 읽는다 — {엑셀 행: 값}.

    병합 칸은 **묶음의 첫 행에만** 값이 있고 나머지는 None 이다. 「항목」은 처음부터
    병합이었고 「질문」은 갈래 행을 합치면서 병합됐다. 그냥 읽으면 둘째 행부터 항목이
    빈 문자열이라 `--item` 이 조용히 그 행들을 버리고, 질문은 빈 채로 에이전트에 들어간다.
    """
    out, cur = {}, ""
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v not in (None, ""):
            cur = str(v)
        out[r] = cur
    return out


def _pins(cell: str) -> list[str]:
    """「대상 고객」 칸에서 kb-pin 을 뽑는다. 여러 개면 대조 행이라 전부 돈다."""
    import re
    return re.findall(r"\d{6}-\d{7}", cell or "")


#: kb-pin → 이름. 대조 행의 답 묶음에 붙이는 이름표가 여기서 나온다. 예전에는 칸의 첫
#: 낱말(`who.split()[0]`)을 썼는데, 그건 **고객이 몇 명이든 언제나 첫 고객 이름**이라
#: 세 사람의 답이 전부 「● 송도윤」으로 적혔다(실측). 답 내용은 맞았고 이름표만 틀려서,
#: 표만 보면 같은 고객에게 세 번 물은 것처럼 읽힌다.
_NAME_OF = {pin: name for name, pin in QM.PIN.items()}


def _llm_down(tr) -> bool:
    """이번 턴이 **LLM 이 죽어서** 끝났나 (tests/debug/reps.py 와 같은 판정).

    LLM 장애는 그래프 안에서 잡혀 「지금은 답변을 만들 수 없어요」라는 **정상 답변**으로
    나온다 — 예외가 안 올라오니 실행기는 성공으로 세고, 그 칸은 채워진 것이 되어 다시는
    안 돌아간다(429 로 죽은 100번이 그렇게 남아 있었다). 답변 글자로 재면 문구가 바뀔 때
    조용히 놓치므로 트레이스로 잰다.
    """
    from tests.debug import trace as TR
    turn = tr.turns[-1] if getattr(tr, "turns", None) else None
    if turn is None:
        return False
    return any(n.delta.get("llm_error") or n.name == TR.LLM_DOWN_NODE for n in turn.nodes)


#: `--rows` 의 구간 구분자. `-` 은 영문 자판, `~` 는 한글 자판에서 손이 먼저 가는 글자이고,
#: `–`(en dash)는 문서·엑셀에서 복사해 붙일 때 딸려 온다. 셋 다 같은 뜻으로 받는다 —
#: 안 받으면 「왜 안 되지」로 한 번 걸리고, 받아서 손해 볼 일이 없다.
_RANGE_SEPS = "-~–—"


def parse_rows(spec: str) -> set[int]:
    """`--rows` 한 칸 → 번호 집합. 낱개와 구간을 섞어 쓸 수 있다.

        17,54,62        낱개
        10-20           구간(양끝 포함)
        10~20 · 10–20   같은 구간(한글 자판 · 문서에서 복사한 en dash)
        1-9,13,25-30    섞어서

    구간은 **양끝을 포함한다** — 「10번부터 20번까지」라고 말할 때 20 을 빼는 사람은 없다.
    거꾸로 쓴 구간(`20-10`)은 바로잡지 않고 **거절한다**: 뒤집어 받아 주면 오타가 조용히
    도는 20행이 되고, 그건 `--redo` 와 만나면 멀쩡한 답을 덮는다.
    """
    want: set[int] = set()
    for chunk in spec.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        cut = next((i for i, ch in enumerate(chunk) if ch in _RANGE_SEPS and i > 0), None)
        if cut is None:
            if not chunk.isdigit():
                raise SystemExit(f"--rows: 번호가 아닙니다 — {chunk!r}")
            want.add(int(chunk))
            continue
        lo, hi = chunk[:cut].strip(), chunk[cut + 1:].strip()
        if not (lo.isdigit() and hi.isdigit()):
            raise SystemExit(f"--rows: 구간이 «숫자-숫자» 가 아닙니다 — {chunk!r}")
        if int(lo) > int(hi):
            raise SystemExit(f"--rows: 구간이 거꾸로입니다 — {chunk!r} (작은 번호를 앞에)")
        want.update(range(int(lo), int(hi) + 1))
    if not want:
        raise SystemExit("--rows: 고른 번호가 없습니다")
    return want


def _select(ws, idx, args) -> list[int]:
    """돌릴 행 번호(엑셀 행 인덱스)를 고른다."""
    want_rows = parse_rows(args.rows) if args.rows else None
    items = _filled(ws, idx["항목"])          # 둘 다 세로 병합이다 — _filled 머리말
    picked = []
    # 번호로 지목했는데 빠진 행은 **왜 빠졌는지** 모아 둔다. 「돌릴 행이 없다」만 남으면
    # 손으로 돌리는 행(137)을 지목한 사람이 실행기가 고장난 줄 안다(2026-09-22 실측).
    skipped: dict = {}
    for r in range(2, ws.max_row + 1):
        num = ws.cell(r, idx["번호"]).value
        item = items[r]
        demo = str(ws.cell(r, idx["시연"]).value or "")
        pre_cell = str(ws.cell(r, idx["선행 질문"]).value or "")
        done = ws.cell(r, idx["실측 답변"]).value

        if want_rows is not None and num not in want_rows:
            continue
        broken = isinstance(done, str) and FAILED in done
        why = None
        if args.grade and not demo.startswith(args.grade):
            why = f"시연 등급이 {args.grade!r} 가 아니다"
        elif args.new and "신규" not in demo:
            why = "신규 행이 아니다"
        elif args.item and args.item not in item:
            why = f"항목이 {args.item!r} 가 아니다"
        elif args.auto and not QM.is_auto(item):
            why = "사람이 읽는 항목이다(--auto 제외)"
        elif args.manual and QM.is_auto(item):
            why = "자동 판정 항목이다(--manual 제외)"
        elif pre_cell.startswith("["):        # 실행 대상이 아닌 행(상황 표기·장애)
            why = f"실행기가 돌리지 않는 행 — {pre_cell.strip('[] ')}"
        # 실패로 끝난 칸은 «채워진» 것이 아니다 — 그냥 두면 건너뛰어져 영영 안 돌고,
        # `--redo` 로 돌리면 멀쩡한 답까지 다시 친다(429 를 부르는 쪽이다).
        # 부분 실패도 실패다 — 대조 행은 고객 여럿을 한 칸에 적으므로 한 명만 죽으면
        # 머리가 아니라 중간에 이 글이 박힌다. startswith 로 재면 그 칸을 영영 못 잡는다.
        elif args.retry_failed and not broken:
            why = "실패로 끝난 칸이 아니다(--retry-failed 제외)"
        elif done and not broken and not args.redo:
            why = "이미 답이 채워져 있다 — 다시 돌리려면 --redo"
        if why:
            if want_rows is not None:
                skipped[num] = why
            continue
        picked.append(r)
        if args.limit and len(picked) >= args.limit:
            break

    for num, why in skipped.items():
        print(f"⚠ {num}번은 건너뛴다 — {why}")

    # 표에 없는 번호를 달라고 했으면 **말한다.** 구간을 받게 되면서(`parse_rows`) 없는
    # 번호를 포함하기가 쉬워졌는데, 그냥 빠지면 화면에는 「돌릴 행이 없다」만 남아
    # «왜 안 돌지»의 답이 어디에도 없다. 표 번호는 164까지다.
    if want_rows is not None:
        have = {ws.cell(r, idx["번호"]).value for r in range(2, ws.max_row + 1)}
        missing = sorted(n for n in want_rows if n not in have)
        if missing:
            shown = ", ".join(str(n) for n in missing[:12])
            print(f"⚠ 표에 없는 번호 {len(missing)}개는 건너뛴다 — {shown}"
                  + (" …" if len(missing) > 12 else ""))
    return picked


def _body(res: dict) -> str:
    """답변 본문 — 화면 장치를 뗀 글.

    `res["answer"]` 에는 답변 뒤에 **화면 장치**가 붙어 있다: 「── 참고한 자료」 표시
    블록 · 「── 이어서 물어보실 수 있어요」 추천질문 · 「… (네 / 아니오)」 제안 문구.
    셀에 그대로 넣으면 표가 안 읽히고, 근거는 어차피 트레이스에서 본다(머리말).

    떼는 규칙을 여기에 다시 쓰지 않는다 — 에이전트 안에 이미 같은 일을 하는 자리가
    있고(`tools/answered._body`, 「지난 답변 다시 써줘」가 쓰는 것), 규칙을 복사하면
    한쪽만 고쳐지는 자리가 생긴다. 추천질문은 `graph.ask` 가 **기록에 넣기 전에** 떼므로
    히스토리 마지막 턴의 답변을 집으면 그 단계까지 같이 해결된다.
    """
    from pension_agent.consult_agent.tools.answered import _body as strip_devices
    turns = res.get("history") or []
    stored = (turns[-1] or {}).get("answer") if turns else None
    return strip_devices(stored or res.get("answer") or "")


def _answer_text(res: dict, args) -> str:
    """칸에 넣을 글. 기본은 답변 본문만 — 그게 이 스크립트의 요점이다."""
    parts = [_body(res)]
    if args.full or args.links:
        for item in res.get("links") or []:
            parts.append(f"[화면] {item['screen']} {item.get('label') or ''} → {item['url']}")
    if args.full or args.sources:
        from pension_agent.consult_agent.tools import source_lines
        lines = [ln for s in res.get("sources") or [] for ln in source_lines(s)]
        if lines:
            parts.append("─ 근거\n" + "\n".join(lines))
    return "\n\n".join(p for p in parts if p)


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(
        prog="python -m scripts.run_question_map",
        description="질문 리스트를 돌려 docs/QUESTION_MAP.xlsx 의 「실측 답변」을 채운다.")
    pick = ap.add_argument_group("돌릴 행 고르기")
    pick.add_argument("--rows", help="번호 목록·구간 (예: 17,54,62 · 10-20 · 1-9,13)")
    pick.add_argument("--grade", choices=["◎", "○", "–"], help="등급으로")
    pick.add_argument("--new", action="store_true",
                      help="기존 대본·QA 에 없던 질문만 (「시연」 칸이 «신규»)")
    pick.add_argument("--item", help="항목 이름 일부 (예: 세액공제)")
    pick.add_argument("--auto", action="store_true",
                      help="답이 대조로 갈리는 항목만 (question_map.AUTO_ITEMS)")
    pick.add_argument("--manual", action="store_true",
                      help="사람이 읽어야 판정되는 항목만 (--auto 의 나머지)")
    pick.add_argument("--limit", type=int, help="앞에서 N개만")
    pick.add_argument("--redo", action="store_true",
                      help="이미 채워진 칸도 다시 돌려 덮어쓴다 (기본은 건너뛴다)")
    pick.add_argument("--retry-failed", action="store_true",
                      help="[실행 실패] 로 끝난 칸만 다시 돌린다")
    rec = ap.add_argument_group("무엇을 칸에 적나 (기본: 답변 본문만)")
    rec.add_argument("--sources", action="store_true", help="근거 목록도 함께")
    rec.add_argument("--links", action="store_true", help="화면 딥링크도 함께")
    rec.add_argument("--full", action="store_true", help="근거 + 화면 딥링크 둘 다")
    run = ap.add_argument_group("실행")
    run.add_argument("--pause", type=float, default=0,
                     help="턴 사이 N초 쉬기 (분당 한도가 있는 키)")
    run.add_argument("--dry-run", action="store_true", help="무엇을 돌릴지만 찍고 끝낸다")
    run.add_argument("--quiet", action="store_true", help="진행 표시를 안 찍는다")
    args = ap.parse_args(argv)

    # .env 를 먼저 읽는다 — 그러지 않으면 아래 머리말이 PENSION_TODAY 를 «고정 안 됨»으로
    # 찍는다(실행 자체는 에이전트가 import 될 때 읽으므로 맞게 돌았다 — 머리말만 거짓말을
    # 했다). 날짜가 틀렸는지를 이 한 줄로 보고 판단하므로 거짓말하면 안 되는 자리다.
    from pension_agent import env as _env
    _env.load()

    wb, ws, idx = _load()
    qs = _filled(ws, idx["질문"])             # 갈래 행은 질문 칸이 세로 병합이다
    picked = _select(ws, idx, args)
    if not picked:
        print("돌릴 행이 없다. 조건을 넓히거나 --redo 를 붙인다.")
        return 0

    today = os.environ.get("PENSION_TODAY")
    print(f"■ {len(picked)}행 · 오늘={today or '실제 날짜(고정 안 됨 — 머리말 ①)'} · "
          f"기록={'답변+근거+화면' if args.full else '답변+근거' if args.sources else '답변+화면' if args.links else '답변 본문만'}")

    if args.dry_run:
        for r in picked:
            n = ws.cell(r, idx["번호"]).value
            q = qs[r]
            who = ws.cell(r, idx["대상 고객"]).value or ""
            pre = ws.cell(r, idx["선행 질문"]).value or ""
            print(f"  {n:>4} [{', '.join(_pins(who)) or '고객 없이'}] {q}"
                  + (f"   ↑ {pre}" if pre else ""))
        return 0

    from tests.debug.runner import session   # 운영 진입점(graph.ask)을 그대로 쓴다

    filled = failed = 0
    for i, r in enumerate(picked, 1):
        num = ws.cell(r, idx["번호"]).value
        q = qs[r]
        who = str(ws.cell(r, idx["대상 고객"]).value or "")
        pre = [x.strip() for x in str(ws.cell(r, idx["선행 질문"]).value or "").split("→") if x.strip()]
        pins = _pins(who) or [None]

        chunks = []
        for cid in pins:
            label = _NAME_OF.get(cid, cid or "") if cid and len(pins) > 1 else ""
            tag = f"[{i}/{len(picked)}] {num} {q[:40]}" + (f" ({label})" if label else "")
            print(f"\n▶ {tag}")
            on_progress = None if args.quiet else (lambda line: print(f"    · {line}"))
            try:
                with session(customer_id=cid, on_progress=on_progress) as (ask, tr):
                    for p in pre:                      # 앞 턴들 — 답은 버린다
                        print(f"    (앞 턴) {p[:40]}")
                        ask(p)
                    res = ask(q)
                    down = _llm_down(tr)
                text = _answer_text(res, args)
            except Exception as e:                     # noqa: BLE001 — 한 행이 죽어도 계속
                failed += 1
                text = f"{FAILED} {type(e).__name__}: {e}"
                print(f"    ✗ {text}")
            else:
                # LLM 이 죽은 턴은 **답이 아니다.** 그래프가 예외를 삼키고 안내 문장으로
                # 돌려주므로 위 except 로는 안 걸린다 — 그대로 두면 채워진 칸이 되어
                # 다시는 안 돌아간다(429 로 죽은 100번이 그렇게 남아 있었다).
                if down:
                    failed += 1
                    text = f"{FAILED} LLM 다운 — {text}"
                    print(f"    ✗ LLM 다운으로 끝난 턴 ({len(text)}자)")
                else:
                    print(f"    ✓ {len(text)}자")
            chunks.append(f"● {label}\n{text}" if label else text)
            if args.pause:
                time.sleep(args.pause)

        ws.cell(r, idx["실측 답변"]).value = "\n\n".join(chunks)
        filled += 1
        wb.save(XLSX)                                  # 턴마다 저장 — 중간에 끊겨도 남는다

    print(f"\n■ {filled}행 기록" + (f" · {failed}건 실패" if failed else "")
          + f" → {XLSX.relative_to(config.REPO_ROOT)}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
